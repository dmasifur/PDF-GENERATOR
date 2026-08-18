"""Read the student spreadsheet into validated :class:`Student` records.

Row-level problems are collected rather than raised: one malformed row should not
abort a whole cohort.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass

from openpyxl import load_workbook

from .models import (
    DEFAULT_COORDINATOR_NAME,
    DEFAULT_COORDINATOR_POSITION,
    DEFAULT_DUTIES,
    PlacementType,
    Student,
)

DATE_FORMAT = "%d-%b-%Y"

# Canonical column name -> the Student field it populates.
COLUMNS = {
    "placement type": "placement_type",
    "participant name": "participant_name",
    "program name": "program_name",
    "host organisation": "host_organisation",
    "host address": "host_address",
    "host abn": "host_abn",
    "host signatory name": "host_signatory_name",
    "host signatory position": "host_signatory_position",
    "supervisor name": "supervisor_name",
    "placement start date": "start_date",
    "placement finish date": "finish_date",
    "attendance hours": "attendance_hours",
    "position title": "position_title",
    "duties": "duties",
    "agreement date": "agreement_date",
    "mid placement date": "mid_placement_date",
    "coordinator name": "coordinator_name",
    "coordinator position": "coordinator_position",
}

REQUIRED = (
    "placement_type",
    "participant_name",
    "program_name",
    "host_organisation",
    "host_address",
    "host_abn",
    "host_signatory_name",
    "supervisor_name",
    "start_date",
    "finish_date",
    "agreement_date",
)

DATE_FIELDS = ("start_date", "finish_date", "agreement_date", "mid_placement_date")

FOE_ALIASES = {
    "foe",
    "found own employment",
    "found own employment (foe)",
    "case 1",
    "self",
    "self placed",
}
PROVIDER_ALIASES = {
    "provider",
    "provider placed",
    "provider-placed",
    "host",
    "case 2",
    "py",
    "placed",
}

DEFAULTS = {
    "duties": DEFAULT_DUTIES,
    "coordinator_name": DEFAULT_COORDINATOR_NAME,
    "coordinator_position": DEFAULT_COORDINATOR_POSITION,
}


@dataclass
class RowError:
    row: int
    reason: str


def _normalise_header(value: object) -> str:
    text = str(value or "").strip().lower()
    for junk in ("/", "_", ".", ":"):
        text = text.replace(junk, " ")
    return " ".join(text.split())


def _clean(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _format_date(value: object, label: str) -> tuple[str, str | None]:

    if value is None or (isinstance(value, str) and not value.strip()):
        return "", None
    if isinstance(value, dt.datetime):
        return value.strftime(DATE_FORMAT), None
    if isinstance(value, dt.date):
        return value.strftime(DATE_FORMAT), None

    text = str(value).strip()
    for fmt in (
        "%d-%b-%Y",
        "%d-%b-%y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d %b %Y",
        "%d/%m/%y",
    ):
        try:
            return dt.datetime.strptime(text, fmt).strftime(DATE_FORMAT), None
        except ValueError:
            continue
    return "", f"{label}: could not read date {text!r}"


def _parse_placement_type(value: object) -> tuple[PlacementType | None, str | None]:
    text = _clean(value).lower()
    if not text:
        return None, "Placement Type: missing"
    if text in FOE_ALIASES:
        return PlacementType.FOE, None
    if text in PROVIDER_ALIASES:
        return PlacementType.PROVIDER, None
    return None, f"Placement Type: unrecognised value {_clean(value)!r}"


def load_students(
    path: str, sheet: str | None = None
) -> tuple[list[Student], list[RowError]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        # Indexing by name can return a chartsheet, which has no rows; going through
        # .worksheets keeps us to sheets we can actually read.
        worksheets = workbook.worksheets
        if sheet is not None:
            by_name = {candidate.title: candidate for candidate in worksheets}
            if sheet not in by_name:
                raise ValueError(
                    f"sheet {sheet!r} not found; available: {sorted(by_name)}"
                )
            worksheet = by_name[sheet]
        elif not worksheets:
            raise ValueError("spreadsheet has no worksheets")
        else:
            worksheet = worksheets[0]

        rows = worksheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ValueError("spreadsheet is empty - no header row") from None

        index: dict[str, int] = {}
        for position, heading in enumerate(header):
            field = COLUMNS.get(_normalise_header(heading))
            if field is not None and field not in index:
                index[field] = position

        missing = [
            name
            for name, field in COLUMNS.items()
            if field in REQUIRED and field not in index
        ]
        if missing:
            raise ValueError(
                "spreadsheet is missing required column(s): "
                + ", ".join(sorted(missing))
            )

        students: list[Student] = []
        errors: list[RowError] = []

        for offset, row in enumerate(rows, start=2):
            if all(cell is None or not str(cell).strip() for cell in row):
                continue

            # A field whose column runs off the end of a short row is absent.
            cells: dict[str, object] = {
                field: row[position]
                for field, position in index.items()
                if position < len(row)
            }

            problems: list[str] = []
            # Every Student field except placement_type, which is not a string.
            values: dict[str, str] = {}

            placement_type, error = _parse_placement_type(cells.get("placement_type"))
            if error:
                problems.append(error)

            for field in COLUMNS.values():
                if field in ("placement_type", *DATE_FIELDS):
                    continue
                text = _clean(cells.get(field))
                values[field] = text or DEFAULTS.get(field, "")

            unparseable: set[str] = set()
            for field in DATE_FIELDS:
                label = field.replace("_", " ").title()
                formatted, error = _format_date(cells.get(field), label)
                if error:
                    problems.append(error)
                    unparseable.add(field)
                values[field] = formatted

            for field in REQUIRED:
                if field in unparseable or field == "placement_type":
                    continue
                if not values.get(field):
                    problems.append(f"{field.replace('_', ' ').title()}: missing")

            if problems or placement_type is None:
                name = _clean(cells.get("participant_name")) or "<unnamed>"
                errors.append(RowError(offset, f"{name} - " + "; ".join(problems)))
                continue

            students.append(Student(placement_type=placement_type, **values))

        return students, errors
    finally:
        workbook.close()
